"""Build bom/PCB_BOM.xlsx and bom/PCB_BOM.csv from an exported KiCad netlist.

Usage:  python make_bom.py build.net [path/to/openpyxl/site-dir]
Rows are grouped by MPN + value + footprint + DNP. Digi-Key part numbers and
prices are carried over from the previous BOM where the MPN is unchanged.
"""
import csv
import os
import re
import sys

from sexp import child, find, parse

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..')
XLSX = os.path.join(ROOT, 'bom', 'PCB_BOM.xlsx')
CSV = os.path.join(ROOT, 'bom', 'PCB_BOM.csv')

REQUIRED = {  # CLAUDE.md section 4 required parts
    'EPC2304', 'ERUC23-3R2K', 'SS110HL-TP', 'GRM21BD72A225KE01L', 'EGXF161ELL111MK25S',
    'SI86S620BC-IS', 'NME1S0505SC', '105313-1206', 'Si8274GB1-IS1',
}
CUSTOM_FP = 'Custom footprint, drawn in Stage 3'


def nat(ref):
    m = re.match(r'([A-Za-z#]+)(\d+)', ref)
    return (m.group(1), int(m.group(2))) if m else (ref, 0)


def components(path):
    tree = parse(open(path, encoding='utf-8').read())
    comps = []
    for comp in find(tree, 'comp'):
        ref = child(comp, 'ref')
        if not ref or ref[1].startswith('#'):
            continue
        props = {}
        for p in comp:
            if isinstance(p, list) and p and p[0] == 'property':
                nm = child(p, 'name')[1]
                val = child(p, 'value')
                props[nm] = val[1] if val else True
        if props.get('exclude_from_bom'):
            continue
        comps.append(dict(ref=ref[1], value=child(comp, 'value')[1],
                          footprint=(child(comp, 'footprint') or [None, ''])[1],
                          mpn=props.get('MPN', ''), mfr=props.get('Manufacturer', ''),
                          lcsc=props.get('LCSC', ''), dnp=bool(props.get('dnp'))))
    return comps


def old_digikey(openpyxl):
    """MPN -> (Digi-Key PN, unit price) from the previous BOM sheet, if present."""
    out = {}
    if not os.path.exists(XLSX):
        return out
    wb = openpyxl.load_workbook(XLSX)
    ws = wb['Original (report BOM)'] if 'Original (report BOM)' in wb.sheetnames else wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    if 'Manufacturer Part Number' not in hdr:
        return out
    i_mpn, i_dk, i_pr = (hdr.index('Manufacturer Part Number'), hdr.index('Digi-Key Part Number 1'),
                         hdr.index('Unit Price 1'))
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[i_mpn]:
            out[str(row[i_mpn]).replace('-', '').upper()] = (row[i_dk], row[i_pr])
    return out


def main(net, pylib=None):
    if pylib:
        sys.path.insert(0, pylib)
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    comps = components(net)
    groups = {}
    for c in comps:
        key = (c['mpn'], c['value'], c['footprint'], c['dnp'])
        groups.setdefault(key, []).append(c)
    dk = old_digikey(openpyxl)
    rows = []
    for (mpn, value, fp, dnp), cs in groups.items():
        refs = sorted((c['ref'] for c in cs), key=nat)
        first = cs[0]
        d, price = dk.get(mpn.replace('-', '').upper(), (None, None))
        notes = []
        if mpn in REQUIRED:
            notes.append('Required part (CLAUDE.md section 4)')
        if fp.startswith('ECE533_Mods:'):
            notes.append(CUSTOM_FP)
        if dnp:
            notes.append('DNP: footprint only')
        rows.append(dict(refs=refs, qty=len(refs), value=value, mpn=mpn, mfr=first['mfr'], dk=d or '',
                         price=float(price) if price not in (None, '') else None, lcsc=first['lcsc'],
                         fp=fp.split(':')[-1], dnp='DNP' if dnp else '', notes='; '.join(notes)))
    rows.sort(key=lambda r: (r['dnp'] != '', nat(r['refs'][0])))

    cols = ['Item', 'Qty', 'Reference designators', 'Value', 'Manufacturer part number', 'Manufacturer',
            'Digi-Key PN', 'Unit price (USD)', 'Ext. price (USD)', 'LCSC PN', 'Footprint', 'DNP', 'Notes']
    # CSV (for diffs)
    with open(CSV, 'w', newline='', encoding='utf-8') as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        for i, r in enumerate(rows, 1):
            w.writerow([i, r['qty'], ','.join(r['refs']), r['value'], r['mpn'], r['mfr'], r['dk'],
                        '' if r['price'] is None else f"{r['price']:.2f}", '', r['lcsc'], r['fp'],
                        r['dnp'], r['notes']])
    # XLSX: new BOM sheet first, previous sheet kept for traceability
    if os.path.exists(XLSX):
        wb = openpyxl.load_workbook(XLSX)
        for ws in list(wb.worksheets):
            if ws.title in ('BOM', 'Sheet1'):
                wb.remove(ws)
        if 'Original (report BOM)' not in wb.sheetnames:
            wb.worksheets[0].title = 'Original (report BOM)'
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    ws = wb.create_sheet('BOM', 0)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='305496')
        c.alignment = Alignment(vertical='center', wrap_text=True)
    for i, r in enumerate(rows, 1):
        n = i + 1
        ws.append([i, r['qty'], ', '.join(r['refs']), r['value'], r['mpn'], r['mfr'], r['dk'],
                   r['price'], f'=IF(H{n}="","",H{n}*B{n})', r['lcsc'], r['fp'], r['dnp'], r['notes']])
        if r['dnp']:
            for c in ws[n]:
                c.font = Font(color='808080', italic=True)
    last = len(rows) + 1
    ws.append([None, f'=SUM(B2:B{last})', 'parts (incl. DNP)', None, None, None, None,
               'Known total:', f'=SUM(I2:I{last})'])
    for c in ws[last + 1]:
        c.font = Font(bold=True)
    widths = [6, 6, 38, 16, 26, 18, 26, 10, 10, 12, 34, 6, 44]
    for col, wdt in zip('ABCDEFGHIJKLM', widths):
        ws.column_dimensions[col].width = wdt
    for row in ws.iter_rows(min_row=2, max_row=last):
        row[2].alignment = Alignment(wrap_text=True, vertical='top')
        row[12].alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:M{last}'
    wb.save(XLSX)
    print(f'{len(rows)} lines, {sum(r["qty"] for r in rows)} parts '
          f'({sum(r["qty"] for r in rows if r["dnp"])} DNP) -> {os.path.normpath(XLSX)}')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
